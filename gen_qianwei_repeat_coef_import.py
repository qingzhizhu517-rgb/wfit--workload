# -*- coding: utf-8 -*-
"""
生成「钱伟·重复系数专项」教学任务导入测试 Excel。

目标教师：钱伟 / 工号 T20270001 / 副教授 / 外聘 / 智能制造学院 (user_id=2003)
学年学期：2026-2027-1（钱伟在该学期无存量数据，计数干净；存量 1 条在 2025-2026-1）

设计目标（按 else/工作量.md:15-18 与 :67-68 的权威口径）：
  C1 为本学期上课重复系数：第一次 1.0，第一次重复(第二次) 0.9，以后 0.8。
  「课程名称一致即为同一门课，不分年级，不以课程代码为准。本专科分别算。」
  G3 的 K 亦为重复系数：第一轮 1，同一教师带不同批次到同一实习单位第二轮起 0.9。
  G2 的 C2 是固定 0.9（不随轮次变化），口径与 G1 不同 —— 一并放进来做对照。

表头必须与 TeachingTaskImportDTO 的 @ExcelProperty 顺序完全一致（17 列）。
其中末两列是本轮新增：
  「班级」    → biz_teaching_task.class_name，附件1「系数说明」据此指名到具体班级；
  「重复次序」→ 显式指定这一行算第几次（1/2/3...）。留空则按同组已入库条数 +1 自动补位，
                分组口径 = 教师 + 学期 + 课程名称 + 授课层次 + 工作量类别（不含课程代码与班级）。

本表除 K 组外一律留空「重复次序」，刻意走自动补位路径；K 组填显式值，验证
「教务可人工指定哪个班算第一次」且显式值优先于文件行序。

Sheet1「教学任务」= 导入数据（EasyExcel 只读第一个 sheet）
Sheet2「用例说明」= 每行的分组、考察点、期望系数与期望工作量、落库校验点
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HEADERS = ["学年学期", "教师工号", "教师姓名", "课程名称", "课程代码", "工作量类别",
           "授课层次", "专业大类", "课程性质", "课程级别", "课程角色", "教学评价",
           "选课人数", "计划学时/天数/周数", "课程系数", "班级", "重复次序"]

SEM = "2026-2027-1"
GH  = "T20270001"
XM  = "钱伟"

# (组, 课程名, 代码, 类别, 层次, 大类, 性质, 级别, 角色, 评价, 人数, 计划值, 系数,
#  班级, 重复次序, 考察点, 期望系数, 期望工作量, 落库校验点)
ROWS = [
    # ---- A 组：本科「高等数学」3 个班。除班级/人数外所有因子刻意保持一致，隔离变量 ----
    ("A", "高等数学", "MATH1001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 45, 64, "",
     "计算机2401", "",
     "同名课第 1 次（自动补位）", "C1=1.0", "70.40", "biz_teaching_task.repeat_order=1 / biz_wl_theory.c1=1.00"),
    ("A", "高等数学", "MATH1001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 52, 64, "",
     "计算机2402", "",
     "同名课第 2 次（第一次重复）", "C1=0.9", "63.36", "repeat_order=2 / c1=0.90"),
    ("A", "高等数学", "MATH1001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 48, 64, "",
     "计算机2403", "",
     "同名课第 3 次及以后", "C1=0.8", "56.32", "repeat_order=3 / c1=0.80"),

    # ---- B 组：专科「高等数学」2 个班。本专科分别算 → 专科重新从第 1 次起算 ----
    ("B", "高等数学", "MATH1002", "G1", "专科", "理工类", "必修", "其他", "独立", "良好", 48, 48, "",
     "机电专2301", "",
     "专科独立计数，第 1 次（不受 A 组 3 次影响）", "C1=1.0", "52.80", "repeat_order=1 / c1=1.00"),
    ("B", "高等数学", "MATH1002", "G1", "专科", "理工类", "必修", "其他", "独立", "良好", 50, 48, "",
     "机电专2302", "",
     "专科第 2 次", "C1=0.9", "47.52", "repeat_order=2 / c1=0.90"),

    # ---- C 组：课程代码与 A 组完全相同但课程名不同 → 按名称判定，不算重复（反向对照）----
    ("C", "概率论", "MATH1001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 40, 32, "",
     "计算机2401", "",
     "代码同 A 组但课程名不同，不得判为重复", "C1=1.0", "35.20", "repeat_order=1 / c1=1.00"),

    # ---- D 组：合堂系数 N 与重复系数叠加，同时覆盖 Q2=1.5（省级一流·主持人）----
    ("D", "大学物理", "PHY2001", "G1", "本科", "理工类", "必修", "省级一流", "主持人", "优秀", 130, 64, "",
     "电气2401+2402合班", "",
     "N=1.1(120~150人) × 第 1 次", "C1=1.0", "116.16", "repeat_order=1 / c1=1.00 / n=1.10"),
    ("D", "大学物理", "PHY2001", "G1", "本科", "理工类", "必修", "省级一流", "主持人", "优秀", 155, 64, "",
     "电气2403+2404+2405合班", "",
     "N=1.2(≥151人) × 第 2 次，两系数叠加", "C1=0.9", "114.05", "repeat_order=2 / c1=0.90 / n=1.20"),

    # ---- E 组：单班不重复，覆盖 Q1=0.8(不合格) + K1=1.0(选修) ----
    ("E", "大学语文", "CHN3001", "G1", "本科", "文史类", "选修", "其他", "独立", "不合格", 60, 32, "",
     "文管2401", "",
     "单班，仅覆盖 Q1=0.8 与 K1=1.0", "C1=1.0", "25.60", "c1=1.00 / k1=1.00 / q1=0.80"),

    # ---- F 组：G2 同名实践课 2 个班。C2 恒 0.9，不随轮次递减（与 G1 口径不同）----
    ("F", "电工实验", "EE4001", "G2", "本科", "理工类", "必修", "其他", "独立", "优秀", 40, 32, 1.0,
     "电气2401", "",
     "G2 第 1 轮", "C2=0.9", "28.80", "biz_wl_practice.c2=0.90（固定值）"),
    ("F", "电工实验", "EE4001", "G2", "本科", "理工类", "必修", "其他", "独立", "优秀", 38, 32, 1.0,
     "电气2402", "",
     "G2 第 2 轮，C2 仍应为 0.9（与 G1 递减口径的对照组）", "C2=0.9", "28.80", "c2=0.90，两行工作量应完全相同"),

    # ---- G 组：G3 同一实习单位 2 轮，K 为重复系数（只有 1.0/0.9 两档，无 0.8）----
    ("G", "生产实习", "PRAC5001", "G3", "本科", "理工类", "必修", "其他", "独立", "优秀", 30, 5, 4.0,
     "机械2201", "",
     "G3 第 1 轮", "K=1.0", "20.00", "repeat_order=1 / biz_wl_internship_training.k=1.00"),
    ("G", "生产实习", "PRAC5001", "G3", "本科", "理工类", "必修", "其他", "独立", "优秀", 28, 5, 4.0,
     "机械2202", "",
     "G3 第 2 轮（同一单位不同批次）", "K=0.9", "18.00", "repeat_order=2 / k=0.90"),

    # ---- H 组：G4 人数触碰上限，检验 CAP_R4_MAX 落库值（已修正为 60）----
    ("H", "综合课程设计", "CD6001", "G4", "本科", "理工类", "必修", "其他", "独立", "优秀", 70, 2, "",
     "智能制造2301-2302", "",
     "R4=70 触碰上限，按 min(R4,60) 计", "上限 60", "48.00", "2 × min(70,60) × 0.4；规则 CAP_R4_MAX=60"),

    # ---- I 组：G5 毕业论文，人数超审批阈值 8 应置 is_over_limit ----
    ("I", "毕业设计", "GRAD7001", "G5", "本科", "理工类", "必修", "其他", "独立", "优秀", 10, 1, 9,
     "智能制造2201", "",
     "R5=10 > APPROVAL_R5_BACHELOR=8", "K5=9", "90.00", "biz_workload_item.is_over_limit=1"),

    # ---- J 组：G6 集中实习，人数触碰上限 20 ----
    ("J", "顶岗实习", "INT8001", "G6", "本科", "理工类", "必修", "其他", "独立", "良好", 25, 4, "",
     "机械2101", "",
     "R6=25 > CAP_R6_MAX=20，超出不计", "上限 20", "32.00", "4 × min(25,20) × 0.4；is_over_limit=1"),

    # ---- K 组：显式「重复次序」。行序在前的填 2、在后的填 1 —— 验证显式值优先于文件行序 ----
    ("K", "数据结构", "DS9001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 45, 48, "",
     "软件2401", 2,
     "显式填 2：虽是文件中该课的第 1 行，仍按第 2 次算", "C1=0.9", "47.52", "repeat_order=2 / c1=0.90"),
    ("K", "数据结构", "DS9001", "G1", "本科", "理工类", "必修", "其他", "独立", "良好", 42, 48, "",
     "软件2402", 1,
     "显式填 1：教务指定此班为第 1 次（不受前一行影响）", "C1=1.0", "52.80", "repeat_order=1 / c1=1.00"),
]

HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# 组内交替底色，便于肉眼区分「同一门课的不同班级」
GROUP_FILL = {"A": "FFF2CC", "B": "E2EFDA", "C": "FCE4D6", "D": "DDEBF7",
              "E": "FFFFFF", "F": "EDEDED", "G": "FFF2CC", "H": "E2EFDA",
              "I": "FCE4D6", "J": "DDEBF7", "K": "E2EFDA"}


def style_head(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER


def build():
    wb = Workbook()

    # ---------- Sheet1：导入数据（EasyExcel 只读第一个 sheet）----------
    ws = wb.active
    ws.title = "教学任务"
    ws.append(HEADERS)
    style_head(ws, len(HEADERS))
    for r in ROWS:
        grp, name, code, gt, edu, major, nature, level, role, ev, cnt, base, coef, cls, order = r[:15]
        ws.append([SEM, GH, XM, name, code, gt, edu, major, nature, level, role, ev,
                   cnt, base, coef, cls, order])
        fill = PatternFill("solid", fgColor=GROUP_FILL[grp])
        for cell in ws[ws.max_row]:
            cell.alignment, cell.border, cell.fill = CENTER, BORDER, fill
    widths = [14, 12, 10, 16, 12, 11, 10, 10, 10, 12, 11, 10, 10, 18, 10, 24, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    # ---------- Sheet2：用例说明（人看的，导入不读）----------
    ws2 = wb.create_sheet("用例说明")
    head2 = ["组", "行号", "课程名称", "班级", "重复次序(填写)", "类别", "层次", "考察点",
             "期望系数", "期望工作量", "落库校验点"]
    ws2.append(head2)
    style_head(ws2, len(head2))
    for idx, r in enumerate(ROWS, start=2):
        grp, name, gt, edu = r[0], r[1], r[3], r[4]
        cls, order = r[13], r[14]
        point, expect_coef, expect_wl, verify = r[15], r[16], r[17], r[18]
        ws2.append([grp, idx, name, cls, order if order != "" else "(留空·自动)",
                    gt, edu, point, expect_coef, expect_wl, verify])
        fill = PatternFill("solid", fgColor=GROUP_FILL[grp])
        for cell in ws2[ws2.max_row]:
            cell.border, cell.fill = BORDER, fill
            cell.alignment = LEFT if cell.column in (8, 11) else CENTER
    for col, w in zip("ABCDEFGHIJK", [5, 6, 16, 24, 14, 8, 8, 44, 12, 12, 44]):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"

    out = "教学任务导入_钱伟_重复系数专项_%d条.xlsx" % len(ROWS)
    wb.save(out)
    print("已生成:", out, "共", len(ROWS), "行")


if __name__ == "__main__":
    build()
